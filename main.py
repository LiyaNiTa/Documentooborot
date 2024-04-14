# Переменная -> Предложение -> Консоль
#name = "Имя"
#print("Ваше имя: ", name, "верно?")
#print(f"Ваше имя {name}, верно?")

#Задача 1: взять предложение из шаблона и заменить на свои данные
# Специальность: 46.02.01 Документационное обеспечение управления и архивоведение, нормативный срок обучения 2 года 10 месяцев
#specialnoct = '46.02.01 Документационное обеспечение управления и архивоведение'
#srok = '2 года 10 месяцев'
#print('Специальность:', specialnoct + ', нормативный срок обучения', srok)

# Таблица (файл csv) -> DataFrame -> Консоль
# устанавливаем библиотеку для работы с таблицами и dataFrame - pandas pip install pandas
import pandas as pd
# Таблицу, которую загрузили в программу, необходимо хранить в переменной типа DataFrame. Принято их называть df
df = pd.read_csv("test.csv")
print(df)

# задача 2: Создать свой файл test.csv, данные из https://randomdatatools.ru/

# Таблица (файл csv) -> DataFrame -> Документ (файл .docx)
# устанавливаем библиотеку для работы с docx pip install python-docx
import docx
# Документ, куда хотим сохранить данные, необходимо хранить в переменной типа docx.
doc = docx.Document()
# Создадим таблицу в документе. shape - это размер
print("размер", df.shape, "количество строк", df.shape[0], "количество столбцов", df.shape[1])
t = doc.add_table(rows=df.shape[0], cols=df.shape[1])
# Заполняем таблицу
print(df.iat[1, 0]) # берет данные из ячейки
for i in range(df.shape[0]): # перебираем все строки
    for j in range(df.shape[1]): # перебираем все столбцы
        cell = df.iat[i, j]
        t.cell(i, j).text = str(cell) # работаем со строками
# Сохраняем документ
doc.save('table 1.docx')
# Сохраняем много документов
for i in range(5):
    name = "table " + str(i) + ".docx"
    doc.save(name)